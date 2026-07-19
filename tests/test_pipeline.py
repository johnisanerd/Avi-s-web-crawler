import json
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import MagicMock, patch

import openpyxl
import pandas as pd

from src.config import Config
from src.constants import STANDARDIZED_COLUMNS, CLAY_COLUMNS
from src.models import RawLead, Lead, ClayLead
from src.mappers import (
    dict_to_raw_lead,
    raw_lead_to_lead,
    lead_to_clay_lead,
    lead_to_standard_dict,
    clay_lead_to_dict
)
from src.normalize import (
    normalize_phone,
    normalize_website,
    normalize_city,
    normalize_state,
    normalize_zip_code,
    normalize_country,
    normalize_rating,
    normalize_review_count,
    normalize_coordinates,
    normalize_lead
)
from src.validators import validate_lead
from src.filters import check_residential_lead
from src.dedupe import deduplicate_leads
from src.processor import process_leads
from src.exporter import LeadExporter
from src.pipeline import LeadGenerationPipeline, PipelineResult
from src.email_extractor import EmailExtractor
from src.enrichment import LeadEnricher
from src.lead_scorer import LeadScorer
from src.business_intelligence import BusinessIntelligence
from src.revenue_intelligence import RevenueIntelligence
from src.technology_intelligence import TechnologyIntelligence
import requests


class TestConfig(unittest.TestCase):
    """Tests configuration parsing, defaults loading, overrides, and validation rules."""

    @patch.dict(os.environ, {"APIFY_API_TOKEN": "env-token"})
    def test_load_config_from_env(self):
        config = Config.load()
        self.assertEqual(config.apify_token, "env-token")
        self.assertEqual(config.max_results_per_query, 50)
        self.assertEqual(config.min_similarity_score, 85)

    def test_load_config_cli_overrides(self):
        overrides = {
            "token": "cli-token",
            "max_results": 20,
            "min_similarity": 90,
            "queries": "agent,office",
            "locations": "Dallas TX"
        }
        config = Config.load(overrides)
        self.assertEqual(config.apify_token, "cli-token")
        self.assertEqual(config.max_results_per_query, 20)
        self.assertEqual(config.min_similarity_score, 90)
        self.assertEqual(config.search_queries, ["agent", "office"])
        self.assertEqual(config.locations, ["Dallas TX"])

    def test_config_validation_missing_token(self):
        with patch.dict(os.environ, {"APIFY_API_TOKEN": ""}):
            with self.assertRaises(ValueError):
                Config.load()


class TestMappers(unittest.TestCase):
    """Tests dict-to-dataclass and dataclass-to-dict mapper translations."""

    def test_dict_to_raw_lead(self):
        data = {
            "title": "Agency A",
            "categoryName": "Real Estate Agency",
            "address": "123 Main St",
            "totalScore": 4.5,
            "reviewsCount": 10,
            "placeId": "ChIJ123",
            "location": {"lat": 25.7617, "lng": -80.1918}
        }
        raw = dict_to_raw_lead(data)
        self.assertEqual(raw.title, "Agency A")
        self.assertEqual(raw.categoryName, "Real Estate Agency")
        self.assertEqual(raw.address, "123 Main St")
        self.assertEqual(raw.totalScore, 4.5)
        self.assertEqual(raw.reviewsCount, 10)
        self.assertEqual(raw.placeId, "ChIJ123")
        self.assertEqual(raw.location, {"lat": 25.7617, "lng": -80.1918})

    def test_raw_lead_to_lead(self):
        raw = RawLead(
            title="Coldwell Banker",
            categoryName="Real Estate Agency",
            address="123 Ocean Dr",
            city="Miami Beach",
            state="FL",
            postalCode="33139",
            phone="305-555-0199",
            website="http://www.coldwell.com",
            url="https://maps.google.com/123",
            totalScore=4.2,
            reviewsCount=15,
            placeId="ChIJ123",
            location={"lat": 25.7906, "lng": -80.1300},
            searchString="real estate in Miami Beach"
        )
        lead = raw_lead_to_lead(raw)
        self.assertEqual(lead.business_name, "Coldwell Banker")
        self.assertEqual(lead.category, "Real Estate Agency")
        self.assertEqual(lead.website, "http://www.coldwell.com")
        self.assertEqual(lead.phone, "305-555-0199")
        self.assertEqual(lead.street, "123 Ocean Dr")
        self.assertEqual(lead.city, "Miami Beach")
        self.assertEqual(lead.state, "FL")
        self.assertEqual(lead.zip_code, "33139")
        self.assertEqual(lead.country, "")
        self.assertEqual(lead.rating, 4.2)
        self.assertEqual(lead.review_count, 15)
        self.assertEqual(lead.latitude, 25.7906)
        self.assertEqual(lead.longitude, -80.1300)
        self.assertEqual(lead.google_maps_url, "https://maps.google.com/123")
        self.assertEqual(lead.google_place_id, "ChIJ123")
        self.assertEqual(lead.source_queries, {"real estate in Miami Beach"})

    def test_lead_to_clay_lead(self):
        lead = Lead(
            business_name="Luxe Real Estate",
            category="Agency",
            website="luxe.com",
            phone="3055550200",
            street="456 Brickell Ave",
            city="Miami",
            state="FL",
            zip_code="33131",
            country="US",
            rating=5.0,
            review_count=1,
            latitude=25.7,
            longitude=-80.2,
            google_maps_url="maps.com/luxe",
            google_place_id="place_luxe",
            source_queries={"q"}
        )
        clay = lead_to_clay_lead(lead)
        self.assertEqual(clay.brokerage, "Luxe Real Estate")
        self.assertEqual(clay.website, "luxe.com")
        self.assertEqual(clay.city, "Miami")
        self.assertEqual(clay.state, "FL")
        self.assertEqual(clay.phone, "3055550200")
        self.assertEqual(clay.google_maps_url, "maps.com/luxe")


class TestNormalize(unittest.TestCase):
    """Tests input data normalization functions."""

    def test_normalize_phone(self):
        self.assertEqual(normalize_phone("+1 (305) 555-0199"), "13055550199")
        self.assertEqual(normalize_phone("305.555.0199"), "3055550199")
        self.assertIsNone(normalize_phone(None))

    def test_normalize_website(self):
        self.assertEqual(normalize_website("https://www.coldwellbanker.com/"), "coldwellbanker.com")
        self.assertEqual(normalize_website("http://realtyonegroup.com/about"), "realtyonegroup.com")
        self.assertEqual(normalize_website("sub.domain.com/path"), "sub.domain.com")
        self.assertIsNone(normalize_website(None))

    def test_normalize_city_state(self):
        self.assertEqual(normalize_city("  miami beach  "), "Miami Beach")
        self.assertEqual(normalize_state("  fl  "), "FL")

    def test_normalize_lead_coordination(self):
        lead = Lead(
            business_name="Unchanged Name LLC",
            category="  realtor  ",
            website="http://www.site.com",
            phone="(305) 555-1234",
            street="  123 main road  ",
            city="miami",
            state="fl",
            zip_code="  33131  ",
            country="us",
            rating=4.5,
            review_count=10,
            latitude="25.76",
            longitude="-80.19",
            google_maps_url="  https://maps.google.com/123  ",
            google_place_id="id123",
            source_queries=set()
        )
        norm = normalize_lead(lead)
        # Business name MUST NOT be altered
        self.assertEqual(norm.business_name, "Unchanged Name LLC")
        self.assertEqual(norm.category, "realtor")
        self.assertEqual(norm.website, "site.com")
        self.assertEqual(norm.phone, "3055551234")
        self.assertEqual(norm.street, "123 main road")
        self.assertEqual(norm.city, "Miami")
        self.assertEqual(norm.state, "FL")
        self.assertEqual(norm.zip_code, "33131")
        self.assertEqual(norm.country, "US")
        self.assertEqual(norm.latitude, 25.76)
        self.assertEqual(norm.longitude, -80.19)


class TestFilters(unittest.TestCase):
    """Tests positive and negative residential lead matching filters."""

    def test_residential_filter_rules(self):
        # 1. Accepted: standard brokerage category
        lead_ok = Lead(
            business_name="Miami Brokers", category="Real Estate Agency",
            website="", phone="", street="", city="", state="", zip_code="", country="",
            rating=0.0, review_count=0, latitude=None, longitude=None,
            google_maps_url="", google_place_id="", source_queries=set()
        )
        res = check_residential_lead(lead_ok)
        self.assertTrue(res.accepted)

        # 2. Excluded: Property management
        lead_pm = Lead(
            business_name="Sunny Property Management", category="Real Estate Agency",
            website="", phone="", street="", city="", state="", zip_code="", country="",
            rating=0.0, review_count=0, latitude=None, longitude=None,
            google_maps_url="", google_place_id="", source_queries=set()
        )
        res = check_residential_lead(lead_pm)
        self.assertFalse(res.accepted)
        self.assertIn("Excluded business name keyword match", res.reason)

        # 3. Commercial brokerage without "residential" is rejected
        lead_comm = Lead(
            business_name="Metro Commercial Realty", category="Commercial Real Estate Agency",
            website="", phone="", street="", city="", state="", zip_code="", country="",
            rating=0.0, review_count=0, latitude=None, longitude=None,
            google_maps_url="", google_place_id="", source_queries=set()
        )
        res = check_residential_lead(lead_comm)
        self.assertFalse(res.accepted)


class TestDeduplication(unittest.TestCase):
    """Tests duplicate lead removal using Google Place ID, Website, Address, and Name matches."""

    def test_deduplication_priorities(self):
        leads = [
            # Base rich lead
            Lead(
                business_name="Coldwell Banker Realty", category="Agency",
                website="coldwell.com", phone="305-555-1111", street="123 Main St",
                city="Miami", state="FL", zip_code="33131", country="US",
                rating=4.5, review_count=20, latitude=None, longitude=None,
                google_maps_url="maps.com/1", google_place_id="place_1",
                source_queries={"queryA"}
            ),
            # Place ID duplicate (less rich)
            Lead(
                business_name="Coldwell Banker", category="Agency",
                website="", phone="", street="",
                city="Miami", state="FL", zip_code="33131", country="US",
                rating=4.0, review_count=1, latitude=None, longitude=None,
                google_maps_url="", google_place_id="place_1",
                source_queries={"queryB"}
            ),
            # Website duplicate
            Lead(
                business_name="Coldwell Banker Miami Office", category="Agency",
                website="coldwell.com", phone="305-555-2222", street="125 Main St",
                city="Miami", state="FL", zip_code="33131", country="US",
                rating=3.0, review_count=2, latitude=None, longitude=None,
                google_maps_url="maps.com/2", google_place_id="place_2",
                source_queries={"queryC"}
            ),
            # Completely unique lead
            Lead(
                business_name="Luxe Living Realty", category="Agency",
                website="luxeliving.com", phone="305-555-9999", street="555 Ocean Dr",
                city="Miami Beach", state="FL", zip_code="33139", country="US",
                rating=4.9, review_count=50, latitude=None, longitude=None,
                google_maps_url="maps.com/3", google_place_id="place_3",
                source_queries={"queryD"}
            )
        ]

        deduped = deduplicate_leads(leads, min_similarity_score=85)
        # Should collapse the three Coldwell entries into one, and keep Luxe Living
        self.assertEqual(len(deduped), 2)
        
        # Verify the richer record's details were kept, and queries merged
        coldwell = next(l for l in deduped if "Coldwell" in l.business_name)
        self.assertEqual(coldwell.rating, 4.5)
        self.assertEqual(coldwell.review_count, 20)
        self.assertEqual(coldwell.source_queries, {"queryA", "queryB", "queryC"})

    def test_deduplication_geo_key_mismatch_crash(self):
        leads = [
            Lead(
                business_name="Realty Office", category="Agency",
                website="realty.com", phone="305-555-1111", street="123 Main St",
                city="Miami", state="FL", zip_code="33131", country="US",
                rating=4.5, review_count=20, latitude=None, longitude=None,
                google_maps_url="maps.com/1", google_place_id="place_1",
                source_queries={"queryA"}
            ),
            # Duplicate by Place ID but has a different ZIP code
            Lead(
                business_name="Realty Office", category="Agency",
                website="realty.com", phone="305-555-1111", street="123 Main St",
                city="Miami", state="FL", zip_code="33139", country="US",
                rating=4.0, review_count=1, latitude=None, longitude=None,
                google_maps_url="maps.com/2", google_place_id="place_1",
                source_queries={"queryB"}
            )
        ]
        # This call should not crash with ValueError
        deduped = deduplicate_leads(leads, min_similarity_score=85)
        self.assertEqual(len(deduped), 1)


class TestProcessor(unittest.TestCase):
    """Tests process_leads pipeline step coordination."""

    def test_process_leads_pipeline(self):
        raw_items = [
            RawLead(
                title="Coldwell Banker", categoryName="Real Estate Agency",
                address="123 Main St", city="Miami", state="MO", postalCode="33131",
                phone="305-555-0100", website="coldwell.com", url="maps.com/1",
                totalScore=4.5, reviewsCount=10, placeId="place1", searchString="query1"
            ),
            RawLead(
                title="Invalid Commercial", categoryName="Commercial Real Estate Agency",
                address="456 Wall St", city="Miami", state="MO", postalCode="33131",
                phone="305-555-0200", website="commercial.com", url="maps.com/2",
                totalScore=4.0, reviewsCount=2, placeId="place2", searchString="query1"
            )
        ]

        res = process_leads(raw_items)
        self.assertEqual(len(res.leads), 1)
        self.assertEqual(res.leads[0].business_name, "Coldwell Banker")
        self.assertEqual(res.filtered_count, 1)


class TestExporter(unittest.TestCase):
    """Tests file export operations under temporary target folders."""

    def test_exporter_creates_files(self):
        config_dict = {
            "token": "dummy",
            "cleaned_csv_name": "clean.csv",
            "cleaned_xlsx_name": "clean.xlsx",
            "clay_csv_name": "clay.csv"
        }

        with tempfile.TemporaryDirectory() as tmp_dir:
            config_dict["output_data_dir"] = Path(tmp_dir)
            config = Config.load(config_dict)

            leads = [
                Lead(
                    business_name="Brokerage A", category="Agency",
                    website="brokeragea.com", phone="305-555-0199", street="123 Ocean Dr",
                    city="Miami", state="FL", zip_code="33139", country="US",
                    rating=4.5, review_count=10, latitude=25.79, longitude=-80.13,
                    google_maps_url="maps.com/a", google_place_id="id_a",
                    source_queries=set()
                )
            ]

            exporter = LeadExporter(config)
            exporter.export_all(leads)

            csv_file = Path(tmp_dir) / "clean.csv"
            xlsx_file = Path(tmp_dir) / "clean.xlsx"
            clay_file = Path(tmp_dir) / "clay.csv"

            # Check if all files were created
            self.assertTrue(csv_file.exists())
            self.assertTrue(xlsx_file.exists())
            self.assertTrue(clay_file.exists())

            # Verify standard CSV content
            csv_df = pd.read_csv(csv_file)
            self.assertEqual(list(csv_df.columns), STANDARDIZED_COLUMNS)
            self.assertEqual(csv_df.iloc[0]["Business Name"], "Brokerage A")

            # Verify Clay CSV content
            clay_df = pd.read_csv(clay_file)
            self.assertEqual(list(clay_df.columns), CLAY_COLUMNS)
            self.assertEqual(clay_df.iloc[0]["Brokerage"], "Brokerage A")

            # Verify Excel formatting structures
            wb = openpyxl.load_workbook(xlsx_file)
            self.assertIn("Brokerage Leads", wb.sheetnames)
            ws = wb["Brokerage Leads"]
            self.assertTrue(ws.views.sheetView[0].showGridLines)
            self.assertEqual(ws.freeze_panes, "A2")


class TestPipeline(unittest.TestCase):
    """Tests the LeadGenerationPipeline using mocked data providers."""

    @patch("src.scraper.ApifyGoogleMapsScraper")
    @patch("requests.Session.get")
    def test_pipeline_execution(self, mock_get, mock_scraper_class):
        mock_get.return_value = None
        # Stub scraper output
        mock_scraper = MagicMock()
        mock_scraper.scrape_leads.return_value = [
            RawLead(
                title="Pipeline Lead A", categoryName="Real Estate Agency",
                address="123 Main St", city="Miami", state="MO", postalCode="33131",
                phone="305-555-0100", website="pipelinea.com", url="maps.com/1",
                totalScore=4.8, reviewsCount=20, placeId="p1", searchString="test"
            )
        ]
        mock_scraper_class.return_value = mock_scraper

        with tempfile.TemporaryDirectory() as tmp_dir:
            tmp_path = Path(tmp_dir)
            overrides = {
                "token": "test-token",
                "locations": "Kansas City MO",
                "data_dir": tmp_path,
                "raw_data_dir": tmp_path / "raw",
                "processed_data_dir": tmp_path / "processed",
                "output_data_dir": tmp_path / "output",
                "logs_dir": tmp_path / "logs",
                "cleaned_csv_name": "clean.csv",
                "cleaned_xlsx_name": "clean.xlsx",
                "clay_csv_name": "clay.csv"
            }
            config = Config.load(overrides)

            pipeline = LeadGenerationPipeline(config, scraper=mock_scraper)
            res = pipeline.run()

            # Verify structured result
            self.assertTrue(res.success)
            self.assertEqual(res.downloaded_count, 1)
            self.assertEqual(res.processed_count, 1)
            self.assertEqual(res.duplicates_removed, 0)
            self.assertEqual(res.exported_count, 1)

            # Verify raw JSON was persisted properly
            raw_files = list((tmp_path / "raw").glob("raw_leads_*.json"))
            self.assertEqual(len(raw_files), 1)
            with open(raw_files[0], "r", encoding="utf-8") as f:
                raw_json = json.load(f)
                self.assertEqual(raw_json[0]["title"], "Pipeline Lead A")

            # Verify export directories exist and have outputs
            self.assertTrue((tmp_path / "output" / "clean.csv").exists())
            self.assertTrue((tmp_path / "output" / "clean.xlsx").exists())
            self.assertTrue((tmp_path / "output" / "clay.csv").exists())


class TestGeographicFilter(unittest.TestCase):
    """Tests the new geographic relevance filter logic in src/filters.py."""

    def test_passes_geographic_filter(self):
        from src.filters import passes_geographic_filter

        # Helper to generate Lead stub
        def make_lead(state: str, country: str) -> Lead:
            return Lead(
                business_name="Brokerage", category="Agency",
                website="", phone="", street="", city="", state=state, zip_code="", country=country,
                rating=0.0, review_count=0, latitude=None, longitude=None,
                google_maps_url="", google_place_id="", source_queries=set()
            )

        us_countries = {"USA", "US", "UNITED STATES", "UNITED STATES OF AMERICA"}

        # 1. Kansas search accepts MO/KS
        mo_ks_states = {"MO", "KS", "MISSOURI", "KANSAS"}
        self.assertTrue(passes_geographic_filter(make_lead("MO", "USA"), allowed_states=mo_ks_states, allowed_countries=us_countries))
        self.assertTrue(passes_geographic_filter(make_lead("KS", "UNITED STATES"), allowed_states=mo_ks_states, allowed_countries=us_countries))
        self.assertTrue(passes_geographic_filter(make_lead("  ks  ", "  united states of america  "), allowed_states=mo_ks_states, allowed_countries=us_countries))
        self.assertTrue(passes_geographic_filter(make_lead("Missouri", "US"), allowed_states=mo_ks_states, allowed_countries=us_countries))

        # 2. Florida search accepts FL
        fl_states = {"FL", "FLORIDA"}
        self.assertTrue(passes_geographic_filter(make_lead("FL", "USA"), allowed_states=fl_states, allowed_countries=us_countries))
        self.assertTrue(passes_geographic_filter(make_lead("Florida", "US"), allowed_states=fl_states, allowed_countries=us_countries))
        self.assertFalse(passes_geographic_filter(make_lead("MO", "USA"), allowed_states=fl_states, allowed_countries=us_countries))

        # 3. Texas search accepts TX
        tx_states = {"TX", "TEXAS"}
        self.assertTrue(passes_geographic_filter(make_lead("TX", "USA"), allowed_states=tx_states, allowed_countries=us_countries))
        self.assertTrue(passes_geographic_filter(make_lead("Texas", "US"), allowed_states=tx_states, allowed_countries=us_countries))
        self.assertFalse(passes_geographic_filter(make_lead("FL", "USA"), allowed_states=tx_states, allowed_countries=us_countries))

        # 4. Empty state/country are allowed (Google Maps sometimes omits them)
        self.assertTrue(passes_geographic_filter(make_lead("", "USA"), allowed_states=mo_ks_states, allowed_countries=us_countries))
        self.assertTrue(passes_geographic_filter(make_lead("KS", ""), allowed_states=mo_ks_states, allowed_countries=us_countries))
        self.assertTrue(passes_geographic_filter(make_lead("", ""), allowed_states=mo_ks_states, allowed_countries=us_countries))

        # 5. Foreign countries are rejected
        self.assertFalse(passes_geographic_filter(make_lead("KS", "Germany"), allowed_states=mo_ks_states, allowed_countries=us_countries))
        self.assertFalse(passes_geographic_filter(make_lead("MO", "PL"), allowed_states=mo_ks_states, allowed_countries=us_countries))


class TestEnrichment(unittest.TestCase):
    """
    Unit tests for EmailExtractor and LeadEnricher using mocked HTTP requests.
    """
    @patch("requests.Session.get")
    def test_homepage_email(self, mock_get):
        # Homepage has email
        mock_response = MagicMock()
        mock_response.status_code = 200
        mock_response.text = "<html><body>Contact us at info@brokerage.com</body></html>"
        mock_get.return_value = mock_response

        extractor = EmailExtractor()
        res = extractor.extract_from_website("http://testsite.com")
        
        self.assertEqual(res.best_email, "info@brokerage.com")
        self.assertEqual(res.emails, ["info@brokerage.com"])
        # Should only visit homepage
        self.assertEqual(mock_get.call_count, 1)

    @patch("requests.Session.get")
    def test_contact_page_email(self, mock_get):
        # Homepage has no email, but contact page has
        mock_home = MagicMock()
        mock_home.status_code = 200
        mock_home.text = "<html><body>No email here <a href='/contact'>Contact Us</a></body></html>"
        
        mock_contact = MagicMock()
        mock_contact.status_code = 200
        mock_contact.text = "<html><body>Email: sales@brokerage.com</body></html>"
        
        mock_get.side_effect = [mock_home, mock_contact] + [None]*8

        extractor = EmailExtractor()
        res = extractor.extract_from_website("http://testsite.com")
        
        self.assertEqual(res.best_email, "sales@brokerage.com")
        self.assertTrue(mock_get.call_count > 1)

    @patch("requests.Session.get")
    def test_owner_email_and_details_detection(self, mock_get):
        # Test detection of owner details: "John Doe - Owner" and jdoe@brokerage.com
        mock_response = MagicMock()
        mock_response.status_code = 200
        mock_response.text = "<html><body>John Doe - Owner (email: jdoe@brokerage.com)</body></html>"
        mock_get.return_value = mock_response

        extractor = EmailExtractor()
        res = extractor.extract_from_website("http://testsite.com")
        
        self.assertEqual(res.owner_name, "John Doe")
        self.assertEqual(res.owner_title, "Owner")
        self.assertEqual(res.owner_email, "jdoe@brokerage.com")

    @patch("requests.Session.get")
    def test_multiple_emails_ranking(self, mock_get):
        # Multiple emails: info@, sales@, broker@, regular@
        mock_response = MagicMock()
        mock_response.status_code = 200
        # broker@ is highest priority, sales@ is priority 2, regular@ is priority 3
        mock_response.text = "<html><body>regular@brokerage.com, sales@brokerage.com, broker@brokerage.com</body></html>"
        mock_get.return_value = mock_response

        extractor = EmailExtractor()
        res = extractor.extract_from_website("http://testsite.com")
        
        # Best email should be ranked by prefix priority (broker@ first)
        self.assertEqual(res.best_email, "broker@brokerage.com")
        self.assertEqual(res.emails[0], "broker@brokerage.com")
        self.assertEqual(res.emails[1], "sales@brokerage.com")
        self.assertEqual(res.emails[2], "regular@brokerage.com")

    @patch("requests.Session.get")
    def test_no_emails_found(self, mock_get):
        mock_response = MagicMock()
        mock_response.status_code = 200
        mock_response.text = "<html><body>Welcome to our website</body></html>"
        mock_get.return_value = mock_response

        extractor = EmailExtractor()
        res = extractor.extract_from_website("http://testsite.com")
        self.assertEqual(res.emails, [])
        self.assertIsNone(res.best_email)

    @patch("requests.Session.get")
    def test_404_error_graceful_handling(self, mock_get):
        mock_response = MagicMock()
        mock_response.status_code = 404
        mock_get.return_value = mock_response

        extractor = EmailExtractor()
        res = extractor.extract_from_website("http://testsite.com")
        self.assertEqual(res.emails, [])

    @patch("requests.Session.get")
    def test_timeout_and_ssl_failure(self, mock_get):
        mock_get.side_effect = requests.Timeout("Connection timed out")

        extractor = EmailExtractor()
        res = extractor.extract_from_website("http://testsite.com")
        self.assertEqual(res.emails, [])

    def test_duplicate_emails_deduplication(self):
        extractor = EmailExtractor()
        raw = ["info@domain.com", "INFO@domain.com", "info@domain.com"]
        ranked = extractor._rank_emails(raw)
        self.assertEqual(ranked, ["info@domain.com"])

    def test_ignored_placeholders(self):
        extractor = EmailExtractor()
        parsed = extractor._parse_emails("test@example.com info@yourcompany.com valid@realty.com")
        self.assertEqual(parsed, ["valid@realty.com"])

    @patch("requests.Session.get")
    def test_lead_enricher_caching_and_enrichment(self, mock_get):
        mock_response = MagicMock()
        mock_response.status_code = 200
        mock_response.text = "<html><body>contact@realty.com</body></html>"
        mock_get.return_value = mock_response

        lead_a = Lead(
            business_name="Brokerage A", category="Agency", website="http://realty.com",
            phone="", street="", city="", state="MO", zip_code="", country="US",
            rating=0.0, review_count=0, latitude=None, longitude=None,
            google_maps_url="", google_place_id="", source_queries=set()
        )
        lead_b = Lead(
            business_name="Brokerage B", category="Agency", website="http://realty.com",
            phone="", street="", city="", state="MO", zip_code="", country="US",
            rating=0.0, review_count=0, latitude=None, longitude=None,
            google_maps_url="", google_place_id="", source_queries=set()
        )

        enricher = LeadEnricher()
        enriched = enricher.enrich_leads([lead_a, lead_b])

        # Verify both leads enriched
        self.assertEqual(enriched[0].business_email, "contact@realty.com")
        self.assertEqual(enriched[1].business_email, "contact@realty.com")
        
        # Should only visit website once (cache hit for lead_b)
        self.assertEqual(mock_get.call_count, 1)


class TestLeadScoring(unittest.TestCase):
    """
    Unit tests for the LeadScorer opportunities scoring engine.
    """

    def setUp(self):
        # Base constructor for clean stubs
        self.scorer = LeadScorer()

    def _make_lead(self, **kwargs) -> Lead:
        defaults = {
            "business_name": "Brokerage",
            "category": "Agency",
            "website": "http://brokerage.com",
            "phone": "305-555-1111",
            "street": "123 Main St",
            "city": "Miami",
            "state": "FL",
            "zip_code": "33131",
            "country": "US",
            "rating": 4.5,
            "review_count": 100,
            "latitude": None,
            "longitude": None,
            "google_maps_url": "",
            "google_place_id": "",
            "source_queries": set(),
            "social_profiles": None,
            "business_email": None,
            "owner_name": None,
            "owner_title": None,
            "owner_email": None
        }
        defaults.update(kwargs)
        return Lead(**defaults)

    def test_perfect_lead_scoring(self):
        # Max out all rules (Website +10, Phone +5, Business email +20, Owner email +25,
        # Owner detected +10, Rating 4.9 +15, 300 reviews +20, Independent +10, Social links +5 = 120 points, clamped to 100)
        lead = self._make_lead(
            business_name="Elite Independent Realty",
            rating=4.9,
            review_count=300,
            business_email="exec@elite.com",
            owner_email="founder@elite.com",
            owner_name="Robert Founder",
            social_profiles={"facebook": "fb.com/elite"}
        )
        scored = self.scorer.score_lead(lead)
        self.assertEqual(scored.lead_score, 100)
        self.assertEqual(scored.lead_tier, "Platinum")
        
        # Verify breakdown explanation lists correct elements
        self.assertIn("Website exists (+10)", scored.score_breakdown)
        self.assertIn("Business email (+20)", scored.score_breakdown)
        self.assertIn("Owner email (+25)", scored.score_breakdown)
        self.assertIn("Owner detected (+10)", scored.score_breakdown)
        self.assertIn("Phone exists (+5)", scored.score_breakdown)
        self.assertIn("Rating 4.9 (+15)", scored.score_breakdown)
        self.assertIn("300 Google reviews (+20)", scored.score_breakdown)
        self.assertIn("Independent brokerage (+10)", scored.score_breakdown)
        self.assertIn("Social links detected (+5)", scored.score_breakdown)

    def test_average_lead_scoring(self):
        # Website +10, Phone +5, Generic email (+20 - 10 = +10), Rating 4.6 (+10), Reviews 75 (+10), Franchise +0 = 45 points
        lead = self._make_lead(
            business_name="Keller Williams Realty",
            rating=4.6,
            review_count=75,
            business_email="office@gmail.com"
        )
        scored = self.scorer.score_lead(lead)
        self.assertEqual(scored.lead_score, 45)
        self.assertEqual(scored.lead_tier, "Bronze")
        self.assertIn("Generic business email (-10)", scored.score_breakdown)
        self.assertIn("National franchise (+0)", scored.score_breakdown)

    def test_no_website_and_no_phone_penalties(self):
        # Missing website (-20), Missing phone (-5), Business email (+20), rating 3.5 (+0), independent (+10) = 5, clamped to 5
        lead = self._make_lead(
            website="",
            phone="",
            business_email="contact@custom.com",
            rating=3.5,
            review_count=0
        )
        scored = self.scorer.score_lead(lead)
        self.assertEqual(scored.lead_score, 5)
        self.assertEqual(scored.lead_tier, "Low Priority")
        self.assertIn("No website (-20)", scored.score_breakdown)
        self.assertIn("Missing phone (-5)", scored.score_breakdown)

    def test_score_clamping_floor(self):
        # Extremely bad case: No website (-20), Missing phone (-5), Franchise (+0), 0 reviews (+0), 0 rating (+0) = -25 -> clamped to 0
        lead = self._make_lead(
            business_name="RE/MAX Dallas",
            website="",
            phone="",
            rating=0.0,
            review_count=0
        )
        scored = self.scorer.score_lead(lead)
        self.assertEqual(scored.lead_score, 0)
        self.assertEqual(scored.lead_tier, "Low Priority")
    def test_rating_and_review_boundaries(self):
        # Rating 4.8+ (+15), Reviews 250+ (+20)
        lead_high = self._make_lead(rating=4.8, review_count=250)
        self.assertEqual(self.scorer.score_lead(lead_high).lead_score, 60)  # Web(10)+Phone(5)+Rating(15)+Reviews(20)+Indep(10)+NoEmail(0) = 60
        
        # Rating 4.5 (+10), Reviews 100 (+15)
        lead_med = self._make_lead(rating=4.5, review_count=100)
        self.assertEqual(self.scorer.score_lead(lead_med).lead_score, 50)  # Web(10)+Phone(5)+Rating(10)+Reviews(15)+Indep(10) = 50

        # Rating 4.0 (+5), Reviews 50 (+10)
        lead_low = self._make_lead(rating=4.0, review_count=50)
        self.assertEqual(self.scorer.score_lead(lead_low).lead_score, 40)  # Web(10)+Phone(5)+Rating(5)+Reviews(10)+Indep(10) = 40

    def test_config_overrides(self):
        # Verify that changing parameters in weights modifies results without altering lead_scorer.py
        import src.constants
        original_weights = src.constants.LEAD_SCORING_WEIGHTS.copy()
        
        try:
            # Set website exists weight to 50
            src.constants.LEAD_SCORING_WEIGHTS["website_exists"] = 50
            
            # Reconstruct scorer to pick up changes
            scorer = LeadScorer()
            lead = self._make_lead(rating=0.0, review_count=0) # Web(50)+Phone(5)+Rating(0)+Reviews(0)+Indep(10) = 65
            scored = scorer.score_lead(lead)
            self.assertEqual(scored.lead_score, 65)
            self.assertEqual(scored.lead_tier, "Silver")
            self.assertIn("Website exists (+50)", scored.score_breakdown)
        finally:
            # Restore original weights for safety
            src.constants.LEAD_SCORING_WEIGHTS = original_weights


class TestBusinessIntelligence(unittest.TestCase):
    """
    Unit tests for the BusinessIntelligence engine.
    """

    def setUp(self):
        self.bi = BusinessIntelligence()
        self.lead = Lead(
            business_name="Test Brokerage",
            category="Agency",
            website="http://testsite.com",
            phone="305-555-1111",
            street="123 Main St",
            city="Miami",
            state="FL",
            zip_code="33131",
            country="US",
            rating=4.5,
            review_count=100,
            latitude=None,
            longitude=None,
            google_maps_url="",
            google_place_id="",
            source_queries=set()
        )

    def test_single_office(self):
        # HTML contains exactly one zip code
        html = "<html><body>Contact us at our office in Miami, FL 33131</body></html>"
        profiled = self.bi.analyze_lead(self.lead, [html])
        self.assertEqual(profiled.office_count, 1)
        self.assertFalse(profiled.has_multiple_offices)

    def test_multiple_offices(self):
        # HTML contains two unique zip codes
        html = "<html><body>Office 1: Miami, FL 33131. Office 2: Tampa, FL 33602</body></html>"
        profiled = self.bi.analyze_lead(self.lead, [html])
        self.assertEqual(profiled.office_count, 2)
        self.assertTrue(profiled.has_multiple_offices)

    def test_agent_directory(self):
        # HTML contains mailto agent emails and agent profile links
        html = """
        <html><body>
            <a href="mailto:john.doe@testsite.com">John Doe</a>
            <a href="mailto:jane.smith@testsite.com">Jane Smith</a>
            <a href="/agents/bob-johnson">Bob Profile</a>
        </body></html>
        """
        profiled = self.bi.analyze_lead(self.lead, [html])
        # Two unique personal emails and one link; max signal is 2
        self.assertEqual(profiled.estimated_agent_count, 2)

    def test_no_team_page(self):
        # HTML contains no agent signals
        html = "<html><body>Welcome to our real estate home!</body></html>"
        profiled = self.bi.analyze_lead(self.lead, [html])
        self.assertIsNone(profiled.estimated_agent_count)

    def test_founded_year(self):
        # HTML contains founded year phrase
        html = "<html><body>We were founded in 1998 to serve you!</body></html>"
        profiled = self.bi.analyze_lead(self.lead, [html])
        self.assertEqual(profiled.years_in_business, 28) # 2026 - 1998

    def test_no_founded_year(self):
        html = "<html><body>Welcome! We are since here to help you.</body></html>"
        profiled = self.bi.analyze_lead(self.lead, [html])
        self.assertIsNone(profiled.years_in_business)

    def test_social_links(self):
        html = """
        <html><body>
            Follow us on <a href="https://www.linkedin.com/company/testrealty">LinkedIn</a>
            and <a href="http://facebook.com/testrealty">Facebook</a>
        </body></html>
        """
        profiled = self.bi.analyze_lead(self.lead, [html])
        self.assertEqual(profiled.linkedin_url, "https://www.linkedin.com/company/testrealty")
        self.assertEqual(profiled.facebook_url, "http://facebook.com/testrealty")
        self.assertEqual(profiled.social_profiles.get("linkedin"), "https://www.linkedin.com/company/testrealty")
        self.assertEqual(profiled.social_profiles.get("facebook"), "http://facebook.com/testrealty")

    def test_no_social_links(self):
        html = "<html><body>No socials here!</body></html>"
        profiled = self.bi.analyze_lead(self.lead, [html])
        self.assertEqual(profiled.social_profiles, {})
        self.assertIsNone(profiled.linkedin_url)

    def test_luxury_detection(self):
        html = "<html><body>We specialize in luxury homes and fine estates.</body></html>"
        profiled = self.bi.analyze_lead(self.lead, [html])
        self.assertTrue(profiled.has_luxury_division)

    def test_commercial_detection(self):
        html = "<html><body>Find commercial real estate and retail spaces.</body></html>"
        profiled = self.bi.analyze_lead(self.lead, [html])
        self.assertTrue(profiled.has_commercial_services)

    def test_mortgage_detection(self):
        html = "<html><body>Ask about our mortgage and home loan options.</body></html>"
        profiled = self.bi.analyze_lead(self.lead, [html])
        self.assertTrue(profiled.has_mortgage_services)

    def test_property_management_detection(self):
        html = "<html><body>Professional property management and leasing.</body></html>"
        profiled = self.bi.analyze_lead(self.lead, [html])
        self.assertTrue(profiled.has_property_management)

    def test_contact_form_detection(self):
        html = "<html><body><form><input name='email'/><button type='submit'>Submit</button></form></body></html>"
        profiled = self.bi.analyze_lead(self.lead, [html])
        self.assertTrue(profiled.has_contact_form)

    def test_confidence_levels(self):
        # Case 1: High confidence (multiple signals: agent, office, years, socials, features)
        html_high = """
        <html><body>
            We were founded in 1990.
            Contact Miami, FL 33131 or Orlando, FL 32801.
            <a href="mailto:a@test.com">Agent A</a>
            <a href="mailto:b@test.com">Agent B</a>
            <a href="https://linkedin.com/company/test">LI</a>
            <a href="https://facebook.com/test">FB</a>
            Specializing in luxury homes, commercial real estate, and property management.
        </body></html>
        """
        self.assertEqual(self.bi.analyze_lead(self.lead, [html_high]).business_intelligence_confidence, "High")

        # Case 2: Medium confidence
        html_med = """
        <html><body>
            We were founded in 2010.
            <a href="https://linkedin.com/company/test">LI</a>
        </body></html>
        """
        self.assertEqual(self.bi.analyze_lead(self.lead, [html_med]).business_intelligence_confidence, "Medium")

        # Case 3: Low confidence
        html_low = "<html><body>Welcome!</body></html>"
        self.assertEqual(self.bi.analyze_lead(self.lead, [html_low]).business_intelligence_confidence, "Low")


class TestRevenueIntelligence(unittest.TestCase):
    """
    Unit tests for the RevenueIntelligence estimation engine.
    """

    def setUp(self):
        self.revenue_intel = RevenueIntelligence()

    def _make_lead(self, **kwargs) -> Lead:
        defaults = {
            "business_name": "Miami Brokerage",
            "category": "Agency",
            "website": "http://brokerage.com",
            "phone": "305-555-1111",
            "street": "123 Main St",
            "city": "Miami",
            "state": "FL",
            "zip_code": "33131",
            "country": "US",
            "rating": 4.5,
            "review_count": 100,
            "latitude": None,
            "longitude": None,
            "google_maps_url": "",
            "google_place_id": "",
            "source_queries": set(),
            "social_profiles": {},
            "business_email": None,
            "owner_name": None,
            "owner_title": None,
            "owner_email": None,
            "estimated_agent_count": 10,
            "office_count": 1,
            "years_in_business": 10,
            "linkedin_url": None,
            "facebook_url": None,
            "instagram_url": None,
            "youtube_url": None,
            "x_url": None,
            "tiktok_url": None,
            "has_blog": False,
            "has_contact_form": False,
            "has_property_search": False,
            "has_mortgage_services": False,
            "has_property_management": False,
            "has_commercial_services": False,
            "has_luxury_division": False,
            "has_careers_page": False,
            "has_multiple_offices": False,
            "business_intelligence_confidence": "Medium"
        }
        defaults.update(kwargs)
        return Lead(**defaults)

    def test_small_brokerage(self):
        # 5 agents * 10 tx/agent = 50 transactions
        # 50 tx * $15,000 GCI/tx = $750,000 GCI
        # Retained: 750,000 * 20% = $150,000 retained base
        # Low limit: 150,000 * 0.8 = $120,000
        # High limit: 150,000 * 1.2 = $180,000
        lead = self._make_lead(estimated_agent_count=5)
        estimated = self.revenue_intel.estimate_lead(lead)
        
        self.assertEqual(estimated.estimated_transactions, 50)
        self.assertEqual(estimated.estimated_gci, 750000.0)
        self.assertEqual(estimated.estimated_revenue_low, 120000.0)
        self.assertEqual(estimated.estimated_revenue_high, 180000.0)
        self.assertIn("Estimated 5 active agents", estimated.revenue_reasoning)
        self.assertIn("Average 10.0 transactions per agent", estimated.revenue_reasoning)
        self.assertIn("Brokerage retained rate: 20%", estimated.revenue_reasoning)

    def test_medium_brokerage(self):
        # 25 agents -> 250 transactions -> $3,750,000 GCI -> $750,000 retained base
        # Low limit: $600,000; High limit: $900,000
        lead = self._make_lead(estimated_agent_count=25)
        estimated = self.revenue_intel.estimate_lead(lead)
        
        self.assertEqual(estimated.estimated_transactions, 250)
        self.assertEqual(estimated.estimated_gci, 3750000.0)
        self.assertEqual(estimated.estimated_revenue_low, 600000.0)
        self.assertEqual(estimated.estimated_revenue_high, 900000.0)

    def test_large_brokerage(self):
        # 100 agents -> 1000 transactions -> $15,000,000 GCI -> $3,000,000 retained base
        # Low limit: $2,400,000; High limit: $3,600,000
        lead = self._make_lead(estimated_agent_count=100)
        estimated = self.revenue_intel.estimate_lead(lead)
        
        self.assertEqual(estimated.estimated_transactions, 1000)
        self.assertEqual(estimated.estimated_gci, 15000000.0)
        self.assertEqual(estimated.estimated_revenue_low, 2400000.0)
        self.assertEqual(estimated.estimated_revenue_high, 3600000.0)

    def test_luxury_brokerage(self):
        # 10 agents -> 100 transactions -> luxury x1.5 GCI per tx ($22,500) -> $2,250,000 GCI -> $450,000 retained base
        # Low: $360,000; High: $540,000
        lead = self._make_lead(estimated_agent_count=10, has_luxury_division=True)
        estimated = self.revenue_intel.estimate_lead(lead)
        
        self.assertEqual(estimated.estimated_gci, 2250000.0)
        self.assertEqual(estimated.estimated_revenue_low, 360000.0)
        self.assertEqual(estimated.estimated_revenue_high, 540000.0)
        self.assertIn("Luxury division detected (GCI multiplier x1.5)", estimated.revenue_reasoning)

    def test_commercial_brokerage(self):
        # 10 agents -> 100 transactions -> commercial x1.3 GCI ($19,500) -> $1,950,000 GCI
        lead = self._make_lead(estimated_agent_count=10, has_commercial_services=True)
        estimated = self.revenue_intel.estimate_lead(lead)
        
        self.assertEqual(estimated.estimated_gci, 1950000.0)
        self.assertIn("Commercial services detected (GCI multiplier x1.3)", estimated.revenue_reasoning)

    def test_property_management_brokerage(self):
        # 10 agents -> 100 transactions -> property management x1.1 GCI ($16,500) -> $1,650,000 GCI
        lead = self._make_lead(estimated_agent_count=10, has_property_management=True)
        estimated = self.revenue_intel.estimate_lead(lead)
        
        self.assertEqual(estimated.estimated_gci, 1650000.0)
        self.assertIn("Property management services detected (GCI multiplier x1.1)", estimated.revenue_reasoning)

    def test_missing_agent_count(self):
        lead = self._make_lead(estimated_agent_count=None)
        estimated = self.revenue_intel.estimate_lead(lead)
        
        self.assertIsNone(estimated.estimated_transactions)
        self.assertIsNone(estimated.estimated_gci)
        self.assertIsNone(estimated.estimated_revenue_low)
        self.assertIsNone(estimated.estimated_revenue_high)
        self.assertEqual(estimated.revenue_confidence, "Low")
        self.assertIn("Unable to estimate revenue: missing agent count", estimated.revenue_reasoning)

    def test_missing_office_count(self):
        # Even with office count missing, points = agent_count(+4) + years(+2) + BI_conf_med(+1) = 7 (Medium confidence)
        lead = self._make_lead(estimated_agent_count=10, office_count=None, years_in_business=10)
        estimated = self.revenue_intel.estimate_lead(lead)
        
        self.assertEqual(estimated.estimated_transactions, 100)
        self.assertEqual(estimated.revenue_confidence, "Medium")

    def test_confidence_tiers(self):
        # Case 1: High confidence: agent(+4) + office(+2) + years(+2) + multiple_offices(+1) + BI_conf_high(+3) = 12 (High)
        lead_high = self._make_lead(
            estimated_agent_count=20,
            office_count=2,
            years_in_business=15,
            has_multiple_offices=True,
            business_intelligence_confidence="High"
        )
        self.assertEqual(self.revenue_intel.estimate_lead(lead_high).revenue_confidence, "High")

        # Case 2: Medium confidence: agent(+4) + office(+2) = 6 (Medium)
        lead_med = self._make_lead(
            estimated_agent_count=10,
            office_count=1,
            years_in_business=None,
            business_intelligence_confidence="Low"
        )
        self.assertEqual(self.revenue_intel.estimate_lead(lead_med).revenue_confidence, "Medium")

        # Case 3: Low confidence: agent(+4) = 4 (Low)
        lead_low = self._make_lead(
            estimated_agent_count=10,
            office_count=None,
            years_in_business=None,
            business_intelligence_confidence="Low"
        )
        self.assertEqual(self.revenue_intel.estimate_lead(lead_low).revenue_confidence, "Low")

    def test_configuration_overrides(self):
        import src.constants
        original_tx_per_agent = src.constants.AVERAGE_TRANSACTIONS_PER_AGENT
        try:
            # Override average transactions per agent to 20
            src.constants.AVERAGE_TRANSACTIONS_PER_AGENT = 20.0
            
            engine = RevenueIntelligence()
            lead = self._make_lead(estimated_agent_count=10) # 10 agents * 20 tx = 200 tx
            estimated = engine.estimate_lead(lead)
            self.assertEqual(estimated.estimated_transactions, 200)
        finally:
            src.constants.AVERAGE_TRANSACTIONS_PER_AGENT = original_tx_per_agent

    def test_model_version_export(self):
        lead = self._make_lead(estimated_agent_count=10)
        estimated = self.revenue_intel.estimate_lead(lead)
        self.assertEqual(estimated.revenue_model_version, "1.0.0")


class TestTechnologyIntelligence(unittest.TestCase):
    """
    Unit tests for the TechnologyIntelligence classification engine.
    """

    def setUp(self):
        self.tech_intel = TechnologyIntelligence()
        self.lead = Lead(
            business_name="Miami Brokerage",
            category="Agency",
            website="http://brokerage.com",
            phone="305-555-1111",
            street="123 Main St",
            city="Miami",
            state="FL",
            zip_code="33131",
            country="US",
            rating=4.5,
            review_count=100,
            latitude=None,
            longitude=None,
            google_maps_url="",
            google_place_id="",
            source_queries=set()
        )

    def test_wordpress_detection(self):
        # Multiple signatures matched -> High confidence
        html = "<html><body><link rel='stylesheet' href='/wp-content/themes/re/' /> Visit /wp-json/wp/v2/posts</body></html>"
        detected = self.tech_intel.detect_lead(self.lead, [html])
        
        self.assertEqual(detected.cms, "WordPress")
        self.assertEqual(detected.website_platform, "WordPress")
        self.assertEqual(detected.technology_confidence, "High")
        self.assertIn("Detected WordPress via signature: 'wp-content'", detected.technology_reasoning)
        self.assertIn("Detected WordPress via signature: 'wp-json'", detected.technology_reasoning)

    def test_wix_detection(self):
        html = "<html><body><script src='wix-code.js'></script></body></html>"
        detected = self.tech_intel.detect_lead(self.lead, [html])
        self.assertEqual(detected.cms, "Wix")
        self.assertEqual(detected.website_platform, "Wix")
        self.assertEqual(detected.technology_confidence, "Medium") # 1 match

    def test_squarespace_detection(self):
        html = "<html><body><script src='https://static1.squarespace.com/static/'></script></body></html>"
        detected = self.tech_intel.detect_lead(self.lead, [html])
        self.assertEqual(detected.cms, "Squarespace")
        self.assertEqual(detected.website_platform, "Squarespace")

    def test_webflow_detection(self):
        html = "<html><body data-wf-page='12345'><link rel='stylesheet' href='webflow.css'/></body></html>"
        detected = self.tech_intel.detect_lead(self.lead, [html])
        self.assertEqual(detected.cms, "Webflow")
        self.assertEqual(detected.website_platform, "Webflow")
        self.assertEqual(detected.technology_confidence, "High")

    def test_google_analytics_gtm_detection(self):
        html = "<html><body><script>gtag('config', 'UA-12345');</script><script src='googletagmanager.com/gtm.js'></script></body></html>"
        detected = self.tech_intel.detect_lead(self.lead, [html])
        self.assertEqual(detected.analytics_tools, "Google Analytics, Google Tag Manager")

    def test_cloudflare_detection(self):
        html = "<html><body><a href='https://cloudflare.com'>Secure</a></body></html>"
        detected = self.tech_intel.detect_lead(self.lead, [html])
        self.assertEqual(detected.cdn_provider, "Cloudflare")

    def test_idx_broker_detection(self):
        html = "<html><body><a href='http://test.idxbroker.com/idx/search'>Search MLS</a></body></html>"
        detected = self.tech_intel.detect_lead(self.lead, [html])
        self.assertEqual(detected.idx_provider, "IDX Broker")

    def test_showcase_idx_detection(self):
        html = "<html><body><div id='sidx-search'></div><script src='showcaseidx.com/loader.js'></script></body></html>"
        detected = self.tech_intel.detect_lead(self.lead, [html])
        self.assertEqual(detected.idx_provider, "Showcase IDX")
        self.assertEqual(detected.technology_confidence, "High")

    def test_kvcore_detection(self):
        html = "<html><body>Welcome to kvcore.com services!</body></html>"
        detected = self.tech_intel.detect_lead(self.lead, [html])
        self.assertEqual(detected.crm_platform, "kvCORE")

    def test_boomtown_detection(self):
        html = "<html><body>Welcome to boomtownre.com system!</body></html>"
        detected = self.tech_intel.detect_lead(self.lead, [html])
        self.assertEqual(detected.crm_platform, "BoomTown")

    def test_follow_up_boss_detection(self):
        html = "<html><body>Follow up boss is closeup followupboss.com!</body></html>"
        detected = self.tech_intel.detect_lead(self.lead, [html])
        self.assertEqual(detected.crm_platform, "Follow Up Boss")

    def test_intercom_detection(self):
        html = "<html><body><script src='widget.intercom.io/widget/123'></script></body></html>"
        detected = self.tech_intel.detect_lead(self.lead, [html])
        self.assertEqual(detected.chat_provider, "Intercom")

    def test_gravity_forms_detection(self):
        html = "<html><body><div id='gforms_wrapper_1'></div></body></html>"
        detected = self.tech_intel.detect_lead(self.lead, [html])
        self.assertEqual(detected.form_provider, "Gravity Forms")

    def test_cookiebot_detection(self):
        html = "<html><body><script id='Cookiebot' src='cookiebot.com/uc.js'></script></body></html>"
        detected = self.tech_intel.detect_lead(self.lead, [html])
        self.assertEqual(detected.consent_manager, "Cookiebot")
        self.assertEqual(detected.technology_confidence, "High")

    def test_unknown_technology(self):
        html = "<html><body>Simple webpage with no tags or scripts.</body></html>"
        detected = self.tech_intel.detect_lead(self.lead, [html])
        self.assertIsNone(detected.cms)
        self.assertIsNone(detected.analytics_tools)
        self.assertEqual(detected.technology_confidence, "Low")
        self.assertIn("No technology signatures detected", detected.technology_reasoning)

    def test_multiple_technologies(self):
        # Wix (cms), Meta Pixel (analytics), HubSpot Forms (form)
        html = """
        <html><body>
            <script src='wix-code.js'></script>
            <script src='fbevents.js'></script>
            <script src='forms.hsforms.com/embed.js'></script>
        </body></html>
        """
        detected = self.tech_intel.detect_lead(self.lead, [html])
        self.assertEqual(detected.cms, "Wix")
        self.assertEqual(detected.analytics_tools, "Meta Pixel")
        self.assertEqual(detected.form_provider, "HubSpot Forms")

    def test_confidence_levels(self):
        # Case 1: High confidence (multiple signatures for WordPress)
        html_high = "<html><body>wp-content and wp-includes</body></html>"
        self.assertEqual(self.tech_intel.detect_lead(self.lead, [html_high]).technology_confidence, "High")

        # Case 2: Medium confidence (single signature for Wix)
        html_med = "<html><body>wix-code</body></html>"
        self.assertEqual(self.tech_intel.detect_lead(self.lead, [html_med]).technology_confidence, "Medium")

        # Case 3: Low confidence (no matches)
        html_low = "<html><body>plain site</body></html>"
        self.assertEqual(self.tech_intel.detect_lead(self.lead, [html_low]).technology_confidence, "Low")

    def test_configuration_overrides(self):
        import src.constants
        original_sigs = src.constants.TECH_SIGNATURES
        try:
            # Inject a custom test technology override
            src.constants.TECH_SIGNATURES = {
                "Awesome CMS": {
                    "category": "cms",
                    "signatures": ["awesome-cms-platform"]
                }
            }
            engine = TechnologyIntelligence()
            html = "<html><body>Powered by awesome-cms-platform</body></html>"
            detected = engine.detect_lead(self.lead, [html])
            self.assertEqual(detected.cms, "Awesome CMS")
        finally:
            src.constants.TECH_SIGNATURES = original_sigs

    def test_model_version(self):
        html = "<html><body>Plain</body></html>"
        detected = self.tech_intel.detect_lead(self.lead, [html])
        self.assertEqual(detected.technology_model_version, "1.0.0")


if __name__ == "__main__":
    unittest.main()
